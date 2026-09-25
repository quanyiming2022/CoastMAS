"""Bounded, non-executable method documents and tabular indicator definitions."""

import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

from .store import Problem

FORMAT = "coastmas.method"
VERSION = "1.0"
COLUMNS = {
    "concept": ("concept", "指标", "指标名称", "科学含义"),
    "unit": ("unit", "单位"),
    "lower": ("lower", "下限", "参考下限"),
    "upper": ("upper", "上限", "参考上限"),
    "positive": ("positive", "方向", "正负方向"),
    "weight": ("weight", "权重"),
}


def parse_method(name, content, mapping=None):
    if len(content) > 8 * 1024**2:
        raise Problem(413, "METHOD_DOCUMENT_SIZE", "方法文件应不超过8MiB")
    suffix = Path(name).suffix.lower()
    provenance = {
        "filename": Path(name).name,
        "sha256": hashlib.sha256(content).hexdigest(),
        "format": suffix.removeprefix("."),
        "size": len(content),
    }
    try:
        if suffix in {".json", ".yaml", ".yml"}:
            document = (
                json.loads(content)
                if suffix == ".json"
                else yaml.safe_load(content.decode("utf-8-sig"))
            )
            if not isinstance(document, dict) or document.get("format") != FORMAT:
                raise Problem(
                    422,
                    "METHOD_STANDARD",
                    "需要带有方法档案标识的JSON/YAML；指标表可用CSV/XLSX导入",
                )
            if document.get("version") != VERSION:
                raise Problem(
                    422,
                    "METHOD_STANDARD_VERSION",
                    "此方法档案版本尚未接入",
                    {"supported": [VERSION], "received": document.get("version")},
                )
            if set(document) - {"format", "version", "definition"}:
                raise Problem(422, "METHOD_DOCUMENT_FIELDS", "存在未识别的顶层字段，不能静默丢弃")
            definition = document["definition"]
            provenance["standard_version"] = VERSION
        elif suffix in {".csv", ".xlsx"}:
            if suffix == ".csv":
                rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
            else:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if (
                        sum(i.file_size for i in archive.infolist()) > 32 * 1024**2
                        or len(archive.infolist()) > 1000
                    ):
                        raise Problem(413, "METHOD_DOCUMENT_SIZE", "工作簿解压后超过读取预算")
                book = load_workbook(
                    io.BytesIO(content), read_only=True, data_only=False, keep_links=False
                )
                try:
                    if len(book.worksheets) != 1:
                        raise Problem(
                            422, "METHOD_WORKSHEET", "指标方案需明确唯一工作表，请单独导出目标表"
                        )
                    sheet = book.active
                    if (sheet.max_row and sheet.max_row > 65) or (
                        sheet.max_column and sheet.max_column > 64
                    ):
                        raise Problem(422, "METHOD_ROWS", "一个方案最多64项指标")
                    rows = []
                    for row in sheet.iter_rows():
                        if any(cell.data_type == "f" for cell in row):
                            raise Problem(
                                422,
                                "METHOD_FORMULA",
                                "指标方案含公式，请提供经确认的数值，不能猜算缓存值",
                            )
                        rows.append([cell.value for cell in row])
                        if len(rows) > 65:
                            raise Problem(422, "METHOD_ROWS", "一个方案最多64项指标")
                finally:
                    book.close()
            if not 2 <= len(rows) <= 65:
                raise Problem(422, "METHOD_ROWS", "需要表头及1至64项指标")
            headers = [str(value or "").strip() for value in rows[0]]
            if len(headers) != len(set(headers)) or "" in headers:
                raise Problem(422, "METHOD_COLUMNS", "指标表表头不能为空或重复")
            chosen = {key: value for key, value in (mapping or {}).items() if value}
            if set(chosen) - set(COLUMNS):
                raise Problem(422, "METHOD_COLUMNS", "列映射包含未知含义")
            for key, aliases in COLUMNS.items():
                candidates = [h for h in headers if h.casefold() in aliases]
                if key not in chosen and len(candidates) == 1:
                    chosen[key] = candidates[0]
            if "concept" not in chosen or not set(chosen.values()).issubset(headers):
                raise Problem(
                    422,
                    "METHOD_COLUMN_MAPPING",
                    "请集中匹配指标表列名",
                    {"columns": headers, "mapping": chosen, "required": ["concept"]},
                )
            if len(set(chosen.values())) != len(chosen):
                raise Problem(422, "METHOD_COLUMNS", "每一科学含义必须对应独立列")
            indicators = []
            for n, row in enumerate(rows[1:], 2):
                if len(row) != len(headers):
                    raise Problem(422, "METHOD_ROW", f"第{n}行列数不符")
                raw = {key: row[headers.index(column)] for key, column in chosen.items()}
                positive = str(raw.get("positive") or "").strip().lower()
                if positive not in {"", "true", "false", "正向", "负向", "1", "0"}:
                    raise Problem(422, "METHOD_DIRECTION", f"第{n}行方向需为正向或负向")
                indicators.append(
                    {
                        "concept": str(raw["concept"] or "").strip(),
                        "unit": str(raw.get("unit") or "").strip(),
                        "lower": float(raw["lower"])
                        if raw.get("lower") not in (None, "")
                        else None,
                        "upper": float(raw["upper"])
                        if raw.get("upper") not in (None, "")
                        else None,
                        "positive": positive in {"true", "正向", "1"} if positive else None,
                        "weight": float(raw["weight"])
                        if raw.get("weight") not in (None, "")
                        else None,
                    }
                )
            definition = {
                "title": Path(name).stem,
                "purpose": "method",
                "basis": "",
                "profiles": ["csv", "csvw", "geotiff", "geojson", "geopackage"],
                "configuration": {
                    "task": "assessment",
                    "method": "weighted",
                    "indicators": indicators,
                },
            }
            provenance["column_mapping"] = chosen
        else:
            raise Problem(
                422, "METHOD_FORMAT", "方案支持JSON、YAML、CSV、XLSX；算法包使用独立模型入口"
            )
        if not isinstance(definition, dict):
            raise Problem(422, "METHOD_DOCUMENT", "方法定义必须为结构化档案")
        # Reject recursive aliases, excessive nesting and non-finite/non-JSON objects.
        json.dumps(definition, allow_nan=False)
        return definition, provenance
    except Problem:
        raise
    except (
        ValueError,
        TypeError,
        KeyError,
        yaml.YAMLError,
        RecursionError,
        zipfile.BadZipFile,
    ) as exc:
        raise Problem(422, "METHOD_DOCUMENT", "方法文件无效；请检查编码、数值和结构") from exc
