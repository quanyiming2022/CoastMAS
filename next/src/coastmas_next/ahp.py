"""Explicit reciprocal judgments, principal eigenvector, published RI convention.

RI profile (1..10): Saaty table reproduced in IJAHP doi:10.13033/ijahp.v15i1.1040,
Table3. Principal-eigenvector basis: doi:10.13033/ijahp.v2i2.87.
Consistency is a diagnostic of judgments, never independent business accuracy.
"""

import csv
import hashlib
import io
import json
import math
import zipfile
from fractions import Fraction
from typing import Annotated, Literal

import numpy as np
from fastapi import APIRouter, File, Form, Request, UploadFile
from openpyxl import load_workbook
from pydantic import Field

from .contracts import Contract
from .store import Problem

RI = [0.0, 0.0, 0.58, 0.90, 1.12, 1.24, 1.32, 1.41, 1.45, 1.49]


class AhpDefinition(Contract):
    method: Literal["ahp"]
    labels: list[str] = Field(min_length=1, max_length=10)
    matrix: list[list[float | None]] = Field(min_length=1, max_length=10)
    consistency_limit: float = Field(gt=0, le=0.1, allow_inf_nan=False)
    source: dict | None = None


def derive_weights(labels, matrix, indicators, limit):
    n = len(labels)
    if (
        not 1 <= n <= 10
        or len(set(labels)) != n
        or len(set(indicators)) != n
        or set(labels) != set(indicators)
    ):
        raise Problem(
            422,
            "AHP_IDENTITIES",
            "矩阵指标必须唯一且与评价指标完整对应",
            {"matrix": labels, "indicators": indicators},
        )
    if not math.isfinite(limit) or not 0 < limit <= 0.1:
        raise Problem(422, "AHP_THRESHOLD", "本RI档案要求明确且不大于0.1的一致性阈值")
    if len(matrix) != n or any(len(row) != n for row in matrix):
        raise Problem(422, "AHP_DIMENSION", "判断矩阵必须与完整指标集合对应")
    missing = [
        [labels[i], labels[j]]
        for i, row in enumerate(matrix)
        for j, value in enumerate(row)
        if value is None
    ]
    if missing:
        raise Problem(422, "AHP_MISSING", f"判断矩阵缺少{len(missing)}项数值", {"pairs": missing})
    value = np.asarray(matrix, dtype="float64")
    if not np.isfinite(value).all() or (value < 1 / 9 - 1e-12).any() or (value > 9 + 1e-12).any():
        raise Problem(422, "AHP_SCALE", "本档案采用1/9至9的有限正比值，不接受零、负值或无穷")
    if not np.allclose(value.diagonal(), 1, rtol=0, atol=1e-10):
        raise Problem(422, "AHP_DIAGONAL", "同一指标与自身比较必须为1")
    if not np.allclose(value * value.T, 1, rtol=0, atol=1e-8):
        pairs = [
            [labels[i], labels[j]]
            for i in range(n)
            for j in range(i + 1, n)
            if abs(value[i, j] * value[j, i] - 1) > 1e-8
        ]
        raise Problem(
            422, "AHP_RECIPROCAL", "判断矩阵倒数关系不成立，请核对原判断", {"pairs": pairs}
        )
    eigenvalues, eigenvectors = np.linalg.eig(value)
    index = int(np.argmax(eigenvalues.real))
    eigenvalue = float(eigenvalues[index].real)
    vector = eigenvectors[:, index].real
    if vector.sum() < 0:
        vector = -vector
    if (vector <= 0).any() or abs(eigenvalues[index].imag) > 1e-10:
        raise Problem(422, "AHP_EIGENVECTOR", "未取得可靠的正主特征向量")
    weights = vector / vector.sum()
    ci = max(0.0, (eigenvalue - n) / (n - 1)) if n > 1 else None
    cr = ci / RI[n - 1] if n >= 3 else None
    order = [labels.index(key) for key in indicators]
    report = {
        "algorithm": "principal_eigenvector",
        "random_index_profile": "saaty_1_10",
        "reference": "https://ijahp.org/index.php/IJAHP/article/download/1040/842",
        "indicators": indicators,
        "weights": weights[order].tolist(),
        "lambda_max": eigenvalue,
        "consistency_index": ci,
        "random_index": RI[n - 1],
        "consistency_ratio": cr,
        "consistency_limit": limit,
        "consistency_applicable": n >= 3,
        "note": "少于3项，不报告通用CR检验通过"
        if n < 3
        else "仅检验判断一致性，不代表独立科学精度",
    }
    if cr is not None and cr > limit:
        raise Problem(
            422, "AHP_INCONSISTENT", "判断矩阵一致性不足，请核对矛盾判断；未修改原矩阵", report
        )
    return report


def parse_matrix(name, content, indicators, limit, aliases=None):
    if len(content) > 8 * 1024**2:
        raise Problem(413, "AHP_SIZE", "判断矩阵文件超过读取预算")
    try:
        if name.lower().endswith(".csv"):
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        elif name.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    sum(i.file_size for i in archive.infolist()) > 32 * 1024**2
                    or len(archive.infolist()) > 1000
                ):
                    raise Problem(413, "AHP_SIZE", "工作簿解压后超过读取预算")
            book = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(book.worksheets) != 1:
                    raise Problem(422, "AHP_SHEET", "请明确唯一矩阵工作表")
                sheet = book.active
                if (sheet.max_row or 0) > 11 or (sheet.max_column or 0) > 11:
                    raise Problem(422, "AHP_SIZE", "本RI档案最多10项指标")
                rows = []
                for row in sheet.iter_rows():
                    if any(c.data_type == "f" for c in row):
                        raise Problem(
                            422, "AHP_FORMULA", "矩阵含公式，请提供确认值，不读取猜测缓存"
                        )
                    rows.append([c.value for c in row])
            finally:
                book.close()
        else:
            raise Problem(422, "AHP_FORMAT", "矩阵支持CSV和XLSX")
        if not 2 <= len(rows) <= 11:
            raise Problem(422, "AHP_DIMENSION", "需要有行列指标名称的矩阵")
        labels = [str(x or "").strip() for x in rows[0][1:]]
        row_labels = [str(row[0] or "").strip() for row in rows[1:]]
        if (
            len(set(labels)) != len(labels)
            or set(row_labels) != set(labels)
            or len(set(row_labels)) != len(row_labels)
        ):
            raise Problem(422, "AHP_IDENTITIES", "行列指标名称缺失、重复或不一致")
        if any(len(row) != len(labels) + 1 for row in rows[1:]):
            raise Problem(422, "AHP_DIMENSION", "矩阵行列不完整")

        def number(raw):
            if raw is None or str(raw).strip() == "":
                return None
            if len(str(raw)) > 64:
                raise ValueError("ratio too long")
            return float(Fraction(str(raw)))

        matrix = [[number(x) for x in rows[1 + row_labels.index(label)][1:]] for label in labels]
        original_labels = labels[:]
        labels = [(aliases or {}).get(label, label) for label in labels]
        definition = AhpDefinition(
            method="ahp",
            labels=labels,
            matrix=matrix,
            consistency_limit=limit,
            source={
                "filename": name,
                "sha256": hashlib.sha256(content).hexdigest(),
                "original_labels": original_labels,
                "mapping": aliases or {},
            },
        )
        report = derive_weights(labels, matrix, indicators, limit)
        return {"definition": definition.model_dump(mode="json"), "report": report}
    except Problem:
        raise
    except (ValueError, TypeError, IndexError, KeyError, zipfile.BadZipFile) as exc:
        raise Problem(422, "AHP_FILE", "矩阵文件的数值、编码或结构无效") from exc


class Preview(Contract):
    definition: AhpDefinition
    indicators: list[str] = Field(min_length=1, max_length=10)


def router(store):
    routes = APIRouter()

    @routes.post("/api/projects/{project}/ahp/preview")
    def preview(project: str, body: Preview, request: Request):
        with store.engine.connect() as c:
            store.permission(c, request.state.actor["id"], project)
        value = body.definition
        return derive_weights(value.labels, value.matrix, body.indicators, value.consistency_limit)

    @routes.post("/api/projects/{project}/ahp/import")
    async def imported(
        project: str,
        request: Request,
        file: Annotated[UploadFile, File()],
        indicators: Annotated[str, Form()],
        consistency_limit: Annotated[float, Form()],
        aliases: Annotated[str | None, Form()] = None,
    ):
        with store.engine.connect() as c:
            store.permission(c, request.state.actor["id"], project, write=True)
        try:
            names = json.loads(indicators)
            mapping = json.loads(aliases) if aliases else {}
            if (
                not isinstance(names, list)
                or not all(isinstance(x, str) for x in names)
                or not isinstance(mapping, dict)
            ):
                raise ValueError("Invalid names")
        except ValueError as exc:
            raise Problem(422, "AHP_IDENTITIES", "需要明确指标名称及列映射") from exc
        return parse_matrix(
            file.filename or "", await file.read(8 * 1024**2 + 1), names, consistency_limit, mapping
        )

    return routes
